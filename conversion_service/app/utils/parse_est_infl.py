"""Parser for est_infl.xlsx — extracts inflation indicators and saves as JSON."""

from __future__ import annotations

import json
import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
)
logger = logging.getLogger(__name__)

TARGET_INDICATORS: list[str] = [
    "Оценка трендовой инфляции (5-летнее скользящее окно)",
    "Оценка трендовой инфляции (3-летнее скользящее окно)",
    "Модифицированный показатель базовой инфляции, рассчитанный на основе метода усечения",
]

INDICATOR_KEYS: list[str] = [
    "Оценка трендовой инфляции (5-летнее скользящее окно)",
    "Оценка трендовой инфляции (3-летнее скользящее окно)",
    "Модифицированный показатель базовой инфляции, рассчитанный на основе метода усечения",
]


def normalize_text(text: str) -> str:
    """Normalize text for fuzzy matching: lowercase, collapse whitespace."""
    if not text:
        return ""
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def find_indicator_row(ws: Worksheet, target: str, max_rows: int = 100) -> int | None:
    """Find row index (1-based) containing the target indicator in column A."""
    normalized_target = normalize_text(target)
    for row_idx in range(1, max_rows + 1):
        cell = ws.cell(row=row_idx, column=1)
        if cell.value is None:
            continue
        cell_text = str(cell.value)
        if normalized_target in normalize_text(cell_text):
            logger.info("Found indicator at row %d: %s", row_idx, cell_text[:80])
            return row_idx
    logger.warning("Indicator not found: %s", target)
    return None


def extract_dates(ws: Worksheet) -> list[datetime]:
    """Extract dates from row 2, starting from column B."""
    dates: list[datetime] = []
    for row in ws.iter_rows(min_row=2, max_row=2, min_col=2, values_only=False):
        for cell in row:
            if isinstance(cell.value, datetime):
                dates.append(cell.value)
            elif cell.value is not None:
                logger.debug("Skipping non-datetime value in row 2 col %d: %s", cell.column, cell.value)
    logger.info("Extracted %d dates from row 2", len(dates))
    return dates


def extract_indicator_data(
    ws: Worksheet,
    row_idx: int,
    num_dates: int,
) -> list[float | None]:
    """Extract numeric values for an indicator from the given row, starting at column B."""
    values: list[float | None] = []
    for col_idx in range(2, 2 + num_dates):
        cell = ws.cell(row=row_idx, column=col_idx)
        val = cell.value
        if val is None:
            values.append(None)
        elif isinstance(val, (int, float)):
            values.append(float(val))
        else:
            logger.debug("Non-numeric value at row %d col %d: %s", row_idx, col_idx, val)
            values.append(None)
    return values


def build_json_structure(
    dates: list[datetime],
    indicator_data: dict[str, list[float | None]],
) -> dict[str, dict[str, dict[str, float]]]:
    """Build nested dict: year -> month -> indicator_key -> value."""
    result: dict[str, dict[str, dict[str, float]]] = {}
    for i, dt in enumerate(dates):
        year_key = str(dt.year)
        month_key = f"{dt.month:02d}"
        if year_key not in result:
            result[year_key] = {}
        if month_key not in result[year_key]:
            result[year_key][month_key] = {}
        for ind_key, values in indicator_data.items():
            val = values[i] if i < len(values) else None
            if val is not None:
                result[year_key][month_key][ind_key] = val
    return result


def parse_est_infl(input_path: Path, output_path: Path) -> None:
    """Main parsing function."""
    logger.info("Opening workbook: %s", input_path)
    wb = openpyxl.load_workbook(input_path, read_only=True, data_only=True)

    try:
        # Use specific sheet name "Лист 1" as required
        ws = wb["Лист 1"]
        logger.info("Using sheet: %s", ws.title)

        dates = extract_dates(ws)
        if not dates:
            raise ValueError("No dates found in row 2")

        indicator_data: dict[str, list[float | None]] = {}
        for target, key in zip(TARGET_INDICATORS, INDICATOR_KEYS):
            row_idx = find_indicator_row(ws, target)
            if row_idx is None:
                logger.warning("Skipping indicator '%s' - not found", key)
                continue
            values = extract_indicator_data(ws, row_idx, len(dates))
            indicator_data[key] = values
            valid_count = sum(1 for v in values if v is not None)
            logger.info("Indicator '%s': %d/%d valid values", key, valid_count, len(values))

        if not indicator_data:
            raise ValueError("No indicators were successfully extracted")

        json_data = build_json_structure(dates, indicator_data)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("w", encoding="utf-8") as f:
            json.dump(json_data, f, ensure_ascii=False, indent=2)

        logger.info("Successfully saved JSON to: %s", output_path)
        logger.info("Years: %d, Total months: %d", len(json_data), sum(len(m) for m in json_data.values()))

    finally:
        wb.close()


if __name__ == "__main__":
    # Use the actual file path that exists
    INPUT_XLSX = Path("est_infl.xlsx")
    OUTPUT_JSON = Path("est_infl_parsed.json")

    if not INPUT_XLSX.exists():
        logger.error("Input file does not exist: %s", INPUT_XLSX)
    else:
        parse_est_infl(INPUT_XLSX, OUTPUT_JSON)